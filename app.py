import streamlit as st
import pandas as pd
import io
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. ตั้งค่าหน้าเว็บ
st.set_page_config(page_title="Stock Material & Daily Monitor", layout="wide")

st.title("📦 สรุปรายการวัสดุและสารเคมีต่ำกว่า Safety Stock (Daily Report)")

# ส่วน Upload ไฟล์ Excel
uploaded_file = st.file_uploader("📂 อัปโหลดไฟล์ Excel (Other_Material_Thailand...)", type=["xlsx", "xls"])

if uploaded_file is not None:
    try:
        # 2. โหลดข้อมูล Stock Material (เริ่มอ่านที่แถวหัวตาราง header=2)
        stock_df = pd.read_excel(uploaded_file, sheet_name="Stock Material", header=2)
        sap_df = pd.read_excel(uploaded_file, sheet_name="SAP_ZRMM0004")
        
        # ตั้งชื่อคอลัมน์แรก (Column A) เป็น 'Type' เพื่อใช้คัดกรอง
        stock_df.rename(columns={stock_df.columns[0]: 'Type'}, inplace=True)
        
        # ตัดช่องว่างชื่อคอลัมน์ทั้งหมด
        stock_df.columns = [str(c).strip() for c in stock_df.columns]
        sap_df.columns = [str(c).strip() for c in sap_df.columns]
        
        # 3. กรองเฉพาะประเภท 'Regular stock' และ 'Pending notification'
        stock_df['Type'] = stock_df['Type'].astype(str).str.strip()
        stock_df = stock_df[stock_df['Type'].isin(['Regular stock', 'Pending notification'])].copy()
        
        # แปลงตัวเลขคอลัมน์ Safety Stock (H) และ Warehouse Stock (I)
        stock_df['Safety Stock'] = pd.to_numeric(stock_df['Safety Stock'], errors='coerce').fillna(0)
        stock_df['Warehouse Stock'] = pd.to_numeric(stock_df['Warehouse Stock'], errors='coerce').fillna(0)
        
        # 4. กรองเฉพาะรายการที่ Warehouse Stock < Safety Stock
        low_stock = stock_df[stock_df['Warehouse Stock'] < stock_df['Safety Stock']].copy()
        
        # 5. เตรียมข้อมูลและจัดกลุ่ม PR / PO จากชีท SAP_ZRMM0004
        sap_df['Mat. Number'] = sap_df['Mat. Number'].astype(str).str.strip()
        sap_df['PR Number'] = sap_df['PR Number'].dropna().astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
        sap_df['PO Number'] = sap_df['PO Number'].dropna().astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
        
        sap_summary = {}
        for mat, grp in sap_df.groupby('Mat. Number'):
            prs = [p for p in grp['PR Number'].dropna().unique() if p and p.lower() != 'nan']
            pos = [p for p in grp['PO Number'].dropna().unique() if p and p.lower() != 'nan']
            sap_summary[mat] = {
                'PR': ", ".join(prs),
                'PO': ", ".join(pos)
            }
            
        # 6. ประกอบข้อมูลรายงานประจำวัน
        report_data = []
        for _, row in low_stock.iterrows():
            mat_type = row.get('Type', '')
            mat_code = str(row.get('Material Code', '')).strip()
            spec = row.get('Specification', '-')
            desc = row.get('รายละเอียด', '-')
            unit = row.get('Unit', '-')
            mat_grp = row.get('Mat.Group', '-')
            safety_stock = row.get('Safety Stock', 0)
            wh_stock = row.get('Warehouse Stock', 0)
            
            # ตรวจสอบ PR ในคอลัมน์ K (PR (1))
            k_val = row.get('PR (1)', 0)
            has_k_pr = pd.notna(k_val) and str(k_val).strip() not in ['0', '', '0.0', 'nan', '-']
            
            # ดึงข้อมูลจาก SAP
            sap_info = sap_summary.get(mat_code, {'PR': '', 'PO': ''})
            pr_val = sap_info['PR']
            po_val = sap_info['PO']
            
            # ตรวจสอบสถานะการเปิด PR
            has_pr = bool(pr_val) or has_k_pr
            remind_to_buy = "Follow PR&PO" if has_pr else "Buy"
                
            report_data.append({
                'Type': mat_type,
                'Material Code': mat_code,
                'Specification': spec,
                'รายละเอียด': desc,
                'Unit': unit,
                'Mat.Group': mat_grp,
                'Safety Stock': round(safety_stock, 2),
                'Warehouse Stock': round(wh_stock, 2),
                'Remind to buy': remind_to_buy,
                'PR': pr_val if pr_val else ("-" if not has_k_pr else f"มีแจ้งเปิดแล้ว (ยอด {k_val})"),
                'PO': po_val if po_val else "-"
            })
            
        final_df = pd.DataFrame(report_data)
        
        # 7. แสดงตัวชี้วัดภาพรวม (Metrics)
        c1, c2, c3 = st.columns(3)
        c1.metric("📌 สต็อกต่ำกว่า Safety (Regular & Pending)", f"{len(final_df)} รายการ")
        c2.metric("🚨 ต้องเปิด PR ด่วน (Buy)", f"{(final_df['Remind to buy'] == 'Buy').sum()} รายการ")
        c3.metric("⏳ ติดตามสถานะ (Follow PR&PO)", f"{(final_df['Remind to buy'] == 'Follow PR&PO').sum()} รายการ")
        
        st.write("---")
        
        # 8. แต่งสีตารางบนหน้าเว็บ Streamlit
        def highlight_status(val):
            if val == 'Buy':
                return 'background-color: #ff4d4d; color: white; font-weight: bold; text-align: center;'
            elif val == 'Follow PR&PO':
                return 'background-color: #ffa600; color: black; font-weight: bold; text-align: center;'
            return ''

        styler = final_df.style
        if hasattr(styler, 'map'):
            styled_df = styler.map(highlight_status, subset=['Remind to buy'])
        else:
            styled_df = styler.applymap(highlight_status, subset=['Remind to buy'])
            
        styled_df = styled_df.format({'Safety Stock': '{:,.2f}', 'Warehouse Stock': '{:,.2f}'})
        st.dataframe(styled_df, use_container_width=True, height=520)
        
        # 9. สร้างไฟล์ Excel พร้อมใส่สีและสไตล์ให้ตรงกับหน้าเว็บ
        excel_buffer = io.BytesIO()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Daily_Low_Stock_Report"

        # หัวตาราง
        headers = list(final_df.columns)
        ws.append(headers)

        # นิยามสไตล์สำหรับ Excel
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        buy_fill = PatternFill(start_color="FF4D4D", end_color="FF4D4D", fill_type="solid")
        buy_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        follow_fill = PatternFill(start_color="FFA600", end_color="FFA600", fill_type="solid")
        follow_font = Font(name="Calibri", size=11, bold=True, color="000000")

        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )

        # ใส่สีหัวตาราง
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # ใส่ข้อมูลและไฮไลต์สีตามค่าใน Remind to buy
        remind_col_idx = headers.index('Remind to buy') + 1

        for r_idx in range(len(final_df)):
            row_num = r_idx + 2
            for c_idx, col_name in enumerate(headers):
                val = final_df.iloc[r_idx, c_idx]
                cell = ws.cell(row=row_num, column=c_idx + 1, value=val)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

                # ฟอร์แมตตัวเลขจุดทศนิยม
                if col_name in ['Safety Stock', 'Warehouse Stock']:
                    cell.number_format = '#,##0.00'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_name in ['Type', 'Unit', 'Mat.Group', 'PR', 'PO']:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # ไฮไลต์สีช่อง Remind to buy
                if c_idx + 1 == remind_col_idx:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if val == 'Buy':
                        cell.fill = buy_fill
                        cell.font = buy_font
                    elif val == 'Follow PR&PO':
                        cell.fill = follow_fill
                        cell.font = follow_font

        # ขยายความกว้างของคอลัมน์อัตโนมัติให้อ่านสบายตา
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        wb.save(excel_buffer)
        excel_buffer.seek(0)
        
        # 10. ปุ่มดาวน์โหลด Daily Report (Excel)
        st.download_button(
            label="📥 ดาวน์โหลดไฟล์ Daily Report ประจำวัน (.xlsx)",
            data=excel_buffer,
            file_name="Daily_Material_Safety_Stock_Report.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        
    except Exception as e:
        st.error(f"⚠️ เกิดข้อผิดพลาดในการประมวลผลไฟล์: {e}")
